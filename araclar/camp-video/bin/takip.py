#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""CAMPVERSE ders edit takip tablosu.

Paylasilan mufredat tablosu salt okunur oldugu icin ona YAZILMAZ.
Kolonlar mufredat sablonuyla birebir ayni, sonuna 3 kolon eklenir;
boylece satirlar paylasilan tabloya oldugu gibi yapistirilabilir.

Kullanim:
    python3 bin/takip.py ekle --vid <loom> --bolum "STEMI" --kitap-no 1 \
        --sure 18 --link "<drive>" --not "opener eklendi, kesim yok"
    python3 bin/takip.py kur          # xlsx + tsv uretir
"""
import argparse
import json
import os
import time

KOK = os.path.expanduser("~/camp-video/takip")
DEFTER = os.path.join(KOK, "takip.json")
XLSX = os.path.join(KOK, "CAMPVERSE-Ders-Edit-Takip.xlsx")
TSV = os.path.join(KOK, "son_satirlar.tsv")

# Mufredat sablonunun birebir kolon sirasi + 3 ek kolon
BASLIK = ["Marka", "Kategori", "Kurs Adı", "Kitap No", "Kitap Adı", "Bölüm No", "Bölüm Adı",
          "Video No", "Video Başlığı", "Tip", "Süre (dk)", "Önizleme (E/H)", "Çekildi (E/H)",
          "Kurgulandı (E/H)", "Yüklendi (E/H)", "Not",
          "Editli Video Linki (Drive)", "Edit Tarihi", "Edit Notu"]

KITAP = {1: "System-Based Clinical Medicine",
         2: "Musculoskeletal, ENT, Eye & Dermatology",
         3: "Women's, Men's & Sexual Health",
         4: "Psychiatry, Paediatrics & Special Populations",
         5: "Emergency Medicine & Radiology"}


def defter_oku():
    if os.path.exists(DEFTER):
        return json.load(open(DEFTER))
    return {"satir": []}


def satir_yap(k):
    return [
        "AUSCAMP", "", "AMC MCQ", k.get("kitap_no") or "", KITAP.get(k.get("kitap_no"), ""),
        k.get("bolum_no") or "", k.get("bolum") or "", k.get("video_no") or 1,
        ("%s Video %d" % (k.get("bolum") or "", k.get("video_no") or 1)).strip(),
        "video", k.get("sure_dk") or "", "H",
        "E",                                   # Çekildi
        "E" if k.get("kurgulandi", True) else "H",
        "E" if k.get("link") else "H",         # Yüklendi
        k.get("mufredat_not") or "",
        k.get("link") or "", k.get("tarih") or "", k.get("edit_notu") or "",
    ]


def kur():
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter

    d = defter_oku()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ders Edit Takip"
    ws.append(BASLIK)
    dolgu = PatternFill("solid", fgColor="06539F")
    for i, _ in enumerate(BASLIK, 1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = dolgu
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for k in d["satir"]:
        ws.append(satir_yap(k))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(BASLIK)), max(1, len(d["satir"]) + 1))
    genislik = [9, 10, 10, 8, 34, 8, 26, 8, 30, 7, 9, 12, 11, 14, 13, 24, 44, 12, 40]
    for i, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    os.makedirs(KOK, exist_ok=True)
    wb.save(XLSX)

    with open(TSV, "w") as f:
        f.write("\t".join(BASLIK) + "\n")
        for k in d["satir"]:
            f.write("\t".join(str(x) for x in satir_yap(k)) + "\n")
    print("Satir: %d" % len(d["satir"]))
    print("->", XLSX)
    print("->", TSV, "(sec, kopyala, paylasilan tabloya yapistir)")


def ekle(a):
    d = defter_oku()
    kayit = {"vid": a.vid, "bolum": a.bolum, "kitap_no": a.kitap_no, "bolum_no": a.bolum_no,
             "video_no": a.video_no, "sure_dk": a.sure, "link": a.link,
             "edit_notu": getattr(a, "not_"), "tarih": a.tarih or time.strftime("%d.%m.%Y"),
             "kurgulandi": True}
    d["satir"] = [s for s in d["satir"] if s.get("vid") != a.vid] + [kayit]
    os.makedirs(KOK, exist_ok=True)
    with open(DEFTER, "w") as f:
        json.dump(d, f, ensure_ascii=False, indent=1)
    print("Deftere islendi:", a.vid[:8], a.bolum)


def main():
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="komut", required=True)
    e = sub.add_parser("ekle")
    e.add_argument("--vid", required=True)
    e.add_argument("--bolum", required=True)
    e.add_argument("--kitap-no", type=int, default=None)
    e.add_argument("--bolum-no", type=int, default=None)
    e.add_argument("--video-no", type=int, default=1)
    e.add_argument("--sure", type=int, default=None)
    e.add_argument("--link", default="")
    e.add_argument("--tarih", default=None)
    e.add_argument("--not", dest="not_", default="")
    sub.add_parser("kur")
    a = ap.parse_args()
    if a.komut == "ekle":
        ekle(a)
    kur()


if __name__ == "__main__":
    main()
