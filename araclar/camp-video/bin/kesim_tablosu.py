#!/usr/bin/env python3
"""Ders başına: grupta önerilen kesim dakikaları ile gerçekte kesilen aralıklar.

Kesilen aralıklar EDL'den okunur (sessizliğe hizalanmış SON sınırlar), gerekçe
model_kesim.json'dan örtüşmeyle eşlenir. Editli videodaki konum = opener + o ana
kadar kalan süre; kullanıcı editli videoda o saniyeye gidip dikişi dinleyebilir.

Kullanım:
    python3 bin/kesim_tablosu.py 63 64 65 ...      # bölüm numaraları
    python3 bin/kesim_tablosu.py --hepsi
"""
import argparse
import glob
import json
import os
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

KOK = os.path.expanduser("~/camp-video")


def ms(s):
    s = max(0.0, s)
    m, sn = divmod(s, 60)
    return "%d:%04.1f" % (m, sn)


def kayitlar():
    k = json.load(open(os.path.join(KOK, "takip/ders_kayitlari.json")))["kayit"]
    return {x["loom_vid"][:8]: x for x in k}


def ders_satirlari(proje, kayit):
    edit = os.path.join(proje, "edit")
    edl = json.load(open(os.path.join(edit, "edl.json")))
    try:
        model = json.load(open(os.path.join(edit, "model_kesim.json")))
    except FileNotFoundError:
        model = {}
    ham_sure = json.load(open(os.path.join(edit, "sessizlik.json")))["toplam_sure"]
    opener = sum(r["end"] - r["start"] for r in edl["ranges"] if r["source"] == "OPENER")
    ders = [r for r in edl["ranges"] if r["source"] == "DERS"]
    mk = model.get("kesimler") or model.get("kesim") or []

    def gerekce(a, b):
        en, iyi = 0, None
        for k in mk:
            ka, kb = k.get("bas", k.get("bas_s")), k.get("bit", k.get("bit_s"))
            o = min(b, kb) - max(a, ka)
            if o > en:
                en, iyi = o, k
        return iyi or {}

    satir = []
    if ders and ders[0]["start"] > 0.05:
        satir.append(dict(tur="baş kırpma", a=0.0, b=ders[0]["start"], editli=opener,
                          gerekce=(model.get("bas_kirpma") or {}).get("gerekce", ""), talimat=""))
    kalan = 0.0
    for i, r in enumerate(ders):
        kalan += r["end"] - r["start"]
        if i + 1 < len(ders):
            a, b = r["end"], ders[i + 1]["start"]
            g = gerekce(a, b)
            satir.append(dict(tur=g.get("tur", ""), a=a, b=b, editli=opener + kalan,
                              gerekce=g.get("gerekce", ""), talimat=g.get("talimat", "")))
    if ders and ham_sure - ders[-1]["end"] > 0.3:
        satir.append(dict(tur="son kırpma", a=ders[-1]["end"], b=ham_sure, editli=opener + kalan,
                          gerekce=(model.get("son_kirpma") or {}).get("gerekce", ""), talimat=""))

    tal = [p["saniye"] for p in (kayit.get("talimat_cozum") or {}).get("noktalar", [])]
    ham_msg = " / ".join(t["metin"] for t in kayit.get("talimat_ham", [])
                         if not t["metin"].startswith("/9j"))
    ham_msg = re.sub(r"https?://\S+", "", ham_msg).strip(" .") or "(talimat yok)"
    # Talimat dakikası hangi kesimle karşılandı: model etiketi, yoksa ±120 sn yakınlık
    karsilama = []
    for t in tal:
        dk = int(t // 60)
        k = next((s for s in satir if str(s["talimat"]) == str(dk)), None)
        if not k:
            k = min((s for s in satir if s["tur"] not in ("baş kırpma", "son kırpma")),
                    key=lambda s: abs((s["a"] + s["b"]) / 2 - t), default=None)
            if k and abs((k["a"] + k["b"]) / 2 - t) > 120:
                k = None
        karsilama.append((dk, k))
    supheli = model.get("supheli") or []

    def editli_konum(t):
        acc = opener
        for r in ders:
            if r["start"] <= t <= r["end"]:
                return acc + (t - r["start"])
            acc += r["end"] - r["start"]
        return None

    for x in supheli:
        x["_editli"] = editli_konum(x.get("bas", 0))
    return dict(satir=satir, tal=tal, ham_msg=ham_msg, karsilama=karsilama, ham_sure=ham_sure,
                editli_sure=edl.get("total_duration_s") or 0, supheli=supheli)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("no", nargs="*", type=int)
    ap.add_argument("--hepsi", action="store_true")
    ap.add_argument("--out", default=os.path.join(KOK, "takip/Kesim-Tablosu.xlsx"))
    a = ap.parse_args()

    dersler = json.load(open(os.path.join(KOK, "takip/dersler.json")))
    kay = kayitlar()
    sec = [d for d in dersler if a.hepsi or d["no"] in a.no]
    sec.sort(key=lambda d: d["no"])

    wb = Workbook()
    oz = wb.active
    oz.title = "Özet"
    ks = wb.create_sheet("Kesimler")
    sp = wb.create_sheet("Kesilmeyen şüpheliler")
    oz.append(["Ders", "Gruptaki öneri (ham mesaj)", "Önerilen dk", "Talimat karşılandı mı",
               "Kesim sayısı", "Kesilen toplam (sn)", "Ham süre", "Editli süre"])
    ks.append(["Ders", "Gruptaki öneri", "Kesilen aralık (HAM kayıt)", "Süre (sn)",
               "Editli videoda dikiş", "Tür", "Talimat dk", "Gerekçe"])
    sp.append(["Ders", "Aralık (HAM)", "Tür", "Not"])
    so = wb.create_sheet("Kullanıcıya sorular", 1)
    so.append(["#", "Ders", "Editli videoda dinle", "Aralık (HAM)", "Konu", "Karar (kes / kalsın)"])
    sira = 0
    for d in sec:
        p = glob.glob(os.path.join(KOK, "projects", "*-" + d["vid"][:8]))
        if not p or not os.path.exists(os.path.join(p[0], "edit/edl.json")):
            oz.append(["CH%02d %s" % (d["no"], d["chapter"]), "", "", "HENÜZ KURGULANMADI"])
            continue
        r = ders_satirlari(p[0], kay.get(d["vid"][:8], {}))
        ad = "CH%02d %s" % (d["no"], d["chapter"])
        ic = [s for s in r["satir"] if s["tur"] not in ("baş kırpma", "son kırpma")]
        kars = ", ".join("%d. dk → %s" % (dk, ("%s–%s" % (ms(k["a"]), ms(k["b"]))) if k else "KARŞILANMADI")
                         for dk, k in r["karsilama"]) or "—"
        oz.append([ad, r["ham_msg"], ", ".join(str(int(t // 60)) for t in r["tal"]) or "—", kars,
                   len(ic), round(sum(s["b"] - s["a"] for s in r["satir"]), 1),
                   ms(r["ham_sure"]), ms(r["editli_sure"])])
        for s in r["satir"]:
            ks.append([ad, r["ham_msg"] if s is r["satir"][0] else "",
                       "%s – %s" % (ms(s["a"]), ms(s["b"])), round(s["b"] - s["a"], 1),
                       ms(s["editli"]), s["tur"], s["talimat"], s["gerekce"]])
        for s in r["supheli"]:
            sp.append([ad, "%s – %s" % (ms(s.get("bas", 0)), ms(s.get("bit", 0))),
                       s.get("tur", ""), s.get("gerekce", "")])
            g = s.get("gerekce", "")
            if "KULLANICIYA" in g or "dinlenmeli" in g.lower():
                sira += 1
                so.append([sira, ad, ms(s["_editli"]) if s["_editli"] is not None else "kesilen bölgede",
                           "%s – %s" % (ms(s.get("bas", 0)), ms(s.get("bit", 0))),
                           re.sub(r";?\s*KULLANICIYA( SORULACAK)?:?", "", g).strip(" ;"), ""])

    ek = os.path.join(KOK, "takip/kullaniciya_ek.json")
    if os.path.exists(ek):
        adlar = {d["no"]: "CH%02d %s" % (d["no"], d["chapter"]) for d in dersler}
        for e in json.load(open(ek)):
            if not (a.hepsi or e["no"] in a.no or e["no"] == 0):
                continue
            sira += 1
            so.append([sira, adlar.get(e["no"], "TÜM DERSLER"), "",
                       ("%s – %s" % (ms(e["bas"]), ms(e["bit"]))) if e["bas"] is not None else "",
                       e["konu"], ""])

    baslik = PatternFill("solid", fgColor="1F3A5F")
    for ws, gen in ((oz, [34, 46, 12, 40, 12, 16, 10, 10]),
                    (so, [5, 34, 14, 22, 90, 18]),
                    (ks, [34, 30, 24, 10, 16, 15, 10, 80]),
                    (sp, [34, 22, 15, 90])):
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = baslik
            c.alignment = Alignment(vertical="center", wrap_text=True)
        for i, w in enumerate(gen, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(vertical="top", wrap_text=True)
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions
    wb.save(a.out)
    print("->", a.out)


if __name__ == "__main__":
    main()
